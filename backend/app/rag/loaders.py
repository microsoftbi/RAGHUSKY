import mimetypes
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as _DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTS = {".pdf", ".md", ".markdown", ".txt", ".html", ".htm", ".docx"}


def detect_mime(path: Path) -> str:
    mime, _ = mimetypes.guess_type(str(path))
    if mime:
        return mime
    return {
        ".md": "text/markdown",
        ".markdown": "text/markdown",
        ".txt": "text/plain",
        ".html": "text/html",
        ".htm": "text/html",
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }.get(path.suffix.lower(), "application/octet-stream")


def is_supported(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_EXTS


def load_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(str(path))
        pages = [(p.extract_text() or "") for p in reader.pages]
        return "\n\n".join(pages)
    if ext == ".docx":
        doc = _DocxDocument(str(path))
        parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)
    raw = path.read_bytes()
    if ext in {".html", ".htm"}:
        text = raw.decode("utf-8", errors="replace")
        soup = BeautifulSoup(text, "lxml")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        return soup.get_text("\n")
    return raw.decode("utf-8", errors="replace")


def normalize(text: str) -> str:
    lines = [line.strip() for line in text.splitlines()]
    cleaned: list[str] = []
    blank = 0
    for ln in lines:
        if not ln:
            blank += 1
            if blank <= 1:
                cleaned.append("")
        else:
            blank = 0
            cleaned.append(ln)
    return "\n".join(cleaned).strip()
